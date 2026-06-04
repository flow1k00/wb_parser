"""Парсер карточек Wildberries.

Читает список артикулов из Excel, по каждому забирает данные с двух
источников WB (card API и basket-CDN) и сохраняет результат в xlsx
и, если настроен доступ, в Google-таблицу.
"""

from __future__ import annotations

import logging
import os
import random
import time
from datetime import datetime
from typing import Any

import openpyxl
import requests
from openpyxl.styles import Font

try:
    import gspread
    from google.oauth2.service_account import Credentials

    HAS_GSPREAD = True
except ImportError:
    HAS_GSPREAD = False

log = logging.getLogger("wb_parser")

INPUT_FILE = "articles.xlsx"
OUTPUT_FILE = "wb_results.xlsx"
CREDENTIALS_FILE = os.getenv("WB_CREDENTIALS", "credentials.json")
SPREADSHEET_ID = os.getenv("WB_SPREADSHEET_ID", "")
SHEET_NAME = os.getenv("WB_SHEET_NAME", "Лист1")

REQUEST_TIMEOUT = 15
DELAY_RANGE = (1.0, 2.5)
BATCH_SIZE = 50
BATCH_PAUSE = 10
MAX_RETRIES = 3
RETRY_DELAY = 5

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept": "application/json, text/plain, */*",
    "Accept-Language": "ru-RU,ru;q=0.9",
    "Referer": "https://www.wildberries.ru/",
    "Origin": "https://www.wildberries.ru",
    "Connection": "keep-alive",
}

# Разные dest = регионы доставки: цена и наличие зависят от региона,
# поэтому перебираем несколько, пока хоть один не ответит.
CARD_ENDPOINTS = [
    "https://card.wb.ru/cards/v2/detail?appType=1&curr=rub&dest=-1257786&spp=30&nm={nm}",
    "https://card.wb.ru/cards/v2/detail?appType=1&curr=rub&dest=-1&spp=30&nm={nm}",
    "https://card.wb.ru/cards/v2/detail?appType=64&curr=rub&dest=-1257786&spp=30&nm={nm}",
    "https://card.wb.ru/cards/v2/detail?appType=1&curr=rub&dest=123585485&spp=30&nm={nm}",
]

COLUMN_HEADERS = [
    "Артикул",
    "Название",
    "Бренд",
    "Цена (руб)",
    "Цена без скидки",
    "Скидка %",
    "Рейтинг",
    "Отзывы",
    "Остаток",
    "Дата обновления",
    "Ссылка",
]


def get_basket_host(vol: int) -> int:
    """Сопоставляет vol (article // 100000) с номером basket-сервера WB.

    Таблица выведена наблюдением за CDN: товары разложены по серверам
    basket-01..25 в зависимости от диапазона vol. WB периодически
    сдвигает границы, поэтому считать её вечной нельзя.
    """
    ranges = [
        (143, 1),
        (287, 2),
        (431, 3),
        (719, 4),
        (1007, 5),
        (1061, 6),
        (1115, 7),
        (1169, 8),
        (1313, 9),
        (1601, 10),
        (1655, 11),
        (1919, 12),
        (2045, 13),
        (2189, 14),
        (2405, 15),
        (2621, 16),
        (2837, 17),
        (3053, 18),
        (3269, 19),
        (3485, 20),
        (3701, 21),
        (3917, 22),
        (4133, 23),
        (4349, 24),
    ]
    for limit, host in ranges:
        if vol <= limit:
            return host
    return 25


def neighbors(host: int) -> list[int]:
    """Сервер-кандидаты в порядке убывания вероятности.

    Сначала сам host, затем два следующих (промахи таблицы обычно в
    сторону более новых товаров на старших серверах) и один предыдущий.
    Несуществующие номера (< 1) отбрасываем сразу.
    """
    return [h for h in (host, host + 1, host + 2, host - 1) if h >= 1]


def get_product_basket(
    session: requests.Session, article: int
) -> dict[str, Any] | None:
    vol = article // 100000
    part = article // 1000
    for host in neighbors(get_basket_host(vol)):
        url = (
            f"https://basket-{host:02d}.wbbasket.ru"
            f"/vol{vol}/part{part}/{article}/info/ru/card.json"
        )
        try:
            response = session.get(url, timeout=REQUEST_TIMEOUT)
            if response.status_code != 200:
                continue
            data = response.json()
        except (requests.RequestException, ValueError) as err:
            log.debug("basket-%02d недоступен для %d: %s", host, article, err)
            continue

        name = data.get("imt_name") or data.get("subj_name", "")
        if name:
            return {
                "_source": f"basket-{host:02d}",
                "_host": host,
                "id": article,
                "name": name,
                "brand": data.get("selling", {}).get("brand_name", ""),
            }
    return None


def get_price_from_basket(
    session: requests.Session, article: int, host: int
) -> dict[str, Any]:
    vol = article // 100000
    part = article // 1000
    for candidate in neighbors(host):
        url = (
            f"https://basket-{candidate:02d}.wbbasket.ru"
            f"/vol{vol}/part{part}/{article}/info/price-history.json"
        )
        try:
            response = session.get(url, timeout=REQUEST_TIMEOUT)
            if response.status_code != 200:
                continue
            history = response.json()
        except (requests.RequestException, ValueError) as err:
            log.debug(
                "price-history недоступна на basket-%02d для %d: %s",
                candidate,
                article,
                err,
            )
            continue

        if isinstance(history, list) and history:
            price = history[-1].get("price", {}).get("RUB", 0)
            if price:
                return {"salePriceU": price, "priceU": price}
    return {}


def get_price_from_card(session: requests.Session, article: int) -> dict[str, Any]:
    for template in CARD_ENDPOINTS:
        url = template.format(nm=article)
        try:
            response = session.get(url, timeout=REQUEST_TIMEOUT)
            if response.status_code != 200:
                continue
            products = response.json().get("data", {}).get("products", [])
        except (requests.RequestException, ValueError) as err:
            log.debug("card API не ответил для %d: %s", article, err)
            continue

        if not products:
            continue
        for product in products:
            if product.get("id") == article:
                return product
        return products[0]
    return {}


def get_product(session: requests.Session, article: int) -> dict[str, Any] | None:
    """Тянет товар, отдавая приоритет card API, с откатом на basket-CDN."""
    card = get_price_from_card(session, article)
    if card.get("name"):
        return {"_source": "card_api", **card}

    basket = get_product_basket(session, article)
    if basket is None:
        return None

    # price-history — основной источник цены; если пусто, переиспользуем
    # уже полученный ответ card, а не дёргаем сеть второй раз.
    price = get_price_from_basket(session, article, basket["_host"]) or card
    basket.update(
        {
            "salePriceU": price.get("salePriceU"),
            "priceU": price.get("priceU"),
            "sale": price.get("sale", 0),
            "rating": price.get("rating", ""),
            "feedbacks": price.get("feedbacks", ""),
            "sizes": price.get("sizes", []),
        }
    )
    return basket


def normalize(product: dict[str, Any], article: int, now: str) -> dict[str, Any]:
    price_raw = product.get("salePriceU") or product.get("priceU") or 0
    price = round(price_raw / 100) if price_raw else 0
    price_original = round((product.get("priceU") or 0) / 100)

    stock = 0
    for size in product.get("sizes", []):
        for item in size.get("stocks", []):
            stock += item.get("qty", 0)

    return {
        "Артикул": article,
        "Название": product.get("name", ""),
        "Бренд": product.get("brand", ""),
        "Цена (руб)": price,
        "Цена без скидки": price_original,
        "Скидка %": product.get("sale", 0),
        "Рейтинг": product.get("rating", ""),
        "Отзывы": product.get("feedbacks", ""),
        "Остаток": stock,
        "Дата обновления": now,
        "Ссылка": f"https://www.wildberries.ru/catalog/{article}/detail.aspx",
    }


def read_articles(path: str) -> list[int]:
    workbook = openpyxl.load_workbook(path, read_only=True)
    sheet = workbook.worksheets[0]
    articles: list[int] = []
    seen: set[int] = set()

    for row_index, row in enumerate(
        sheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        raw = row[0]
        if raw is None or raw == "":
            continue
        try:
            # Excel хранит целые как float ("12345678.0"), поэтому режем хвост.
            article = int(str(raw).strip().split(".")[0])
        except ValueError:
            log.warning("строка %d: %r не похоже на артикул, пропускаю", row_index, raw)
            continue
        if article not in seen:
            seen.add(article)
            articles.append(article)

    workbook.close()
    return articles


def save_excel(items: list[dict[str, Any]], path: str) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Товары"
    sheet.append(COLUMN_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for item in items:
        sheet.append([item.get(header, "") for header in COLUMN_HEADERS])

    widths = [12, 45, 25, 12, 14, 10, 10, 10, 10, 18, 55]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width

    workbook.save(path)
    log.info("Excel: записано %d строк в %s", len(items), path)


def save_sheets(items: list[dict[str, Any]]) -> None:
    if not HAS_GSPREAD:
        log.warning("Google Sheets пропущен: установите gspread и google-auth")
        return
    if not SPREADSHEET_ID:
        log.warning("Google Sheets пропущен: не задан WB_SPREADSHEET_ID")
        return

    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive",
    ]
    try:
        creds = Credentials.from_service_account_file(CREDENTIALS_FILE, scopes=scopes)
        client = gspread.authorize(creds)
        spreadsheet = client.open_by_key(SPREADSHEET_ID)
        try:
            sheet = spreadsheet.worksheet(SHEET_NAME)
        except gspread.WorksheetNotFound:
            sheet = spreadsheet.add_worksheet(title=SHEET_NAME, rows=2000, cols=20)

        sheet.clear()
        rows = [COLUMN_HEADERS]
        rows.extend(
            [str(item.get(header, "")) for header in COLUMN_HEADERS] for item in items
        )
        sheet.update(rows, value_input_option="RAW")
        log.info("Google Sheets: записано %d строк на лист %s", len(items), SHEET_NAME)
    except FileNotFoundError:
        log.error("Google Sheets: файл %s не найден", CREDENTIALS_FILE)
    except Exception as err:
        log.error("Google Sheets: %s", err)


def fetch_all(
    session: requests.Session, articles: list[int], now: str
) -> dict[int, dict[str, Any]]:
    """Основной проход по артикулам с задержками и паузами между батчами."""
    results: dict[int, dict[str, Any]] = {}
    missed: set[int] = set()
    total = len(articles)

    try:
        for position, article in enumerate(articles, start=1):
            product = get_product(session, article)
            if product:
                row = normalize(product, article, now)
                results[article] = row
                log.info(
                    "[%d/%d] %d — OK (%s) %s, %s ₽",
                    position,
                    total,
                    article,
                    product.get("_source", ""),
                    row["Название"][:40],
                    row["Цена (руб)"],
                )
            else:
                missed.add(article)
                log.info("[%d/%d] %d — не найден", position, total, article)

            if position < total:
                time.sleep(random.uniform(*DELAY_RANGE))
            if position % BATCH_SIZE == 0 and position < total:
                log.info("Пауза %d сек после %d артикулов", BATCH_PAUSE, position)
                time.sleep(BATCH_PAUSE)
    except KeyboardInterrupt:
        log.warning("Прервано вручную, сохраняю собранное")

    retry_missing(session, results, missed, now)
    return results


def retry_missing(
    session: requests.Session,
    results: dict[int, dict[str, Any]],
    missed: set[int],
    now: str,
) -> None:
    """Добивает ненайденные артикулы несколькими заходами."""
    for attempt in range(1, MAX_RETRIES + 1):
        if not missed:
            return
        log.info("Повтор %d/%d: осталось %d", attempt, MAX_RETRIES, len(missed))
        time.sleep(RETRY_DELAY)
        for article in list(missed):
            product = get_product(session, article)
            if product:
                results[article] = normalize(product, article, now)
                missed.discard(article)
            time.sleep(random.uniform(*DELAY_RANGE))


def build_rows(
    articles: list[int], results: dict[int, dict[str, Any]], now: str
) -> list[dict[str, Any]]:
    """Собирает итог в исходном порядке, ненайденным ставит заглушку."""
    rows = []
    for article in articles:
        if article in results:
            rows.append(results[article])
        else:
            rows.append(
                {
                    "Артикул": article,
                    "Название": "Не найден",
                    "Дата обновления": now,
                    "Ссылка": f"https://www.wildberries.ru/catalog/{article}/detail.aspx",
                }
            )
    return rows


def main() -> None:
    articles = read_articles(INPUT_FILE)
    log.info("Артикулов к обработке: %d", len(articles))
    if not articles:
        log.error("Список пуст — нечего парсить")
        return

    session = requests.Session()
    session.headers.update(HEADERS)
    try:
        session.get("https://www.wildberries.ru", timeout=REQUEST_TIMEOUT)
    except requests.RequestException:
        log.warning("Не удалось прогреть сессию, продолжаю без куки")

    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    results = fetch_all(session, articles, now)
    rows = build_rows(articles, results, now)

    save_excel(rows, OUTPUT_FILE)
    save_sheets(rows)
    log.info("Готово. Найдено %d из %d", len(results), len(articles))


if __name__ == "__main__":
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s  %(levelname)-7s  %(message)s",
        datefmt="%H:%M:%S",
    )
    main()
